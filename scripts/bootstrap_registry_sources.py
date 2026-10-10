#!/usr/bin/env python3
"""Merge exact-ID-linked historical workbook evidence into an existing snapshot.

This is a one-time local bootstrap. Monthly captures keep these observations by
merging with the deployed snapshot. Workbooks are read without being modified.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from build_registry_snapshot import MAX_LOCAL_BYTES, json_bytes, read_json, validate_snapshot, write_atomic
from build_seed import CADASTRAL_ID, ROOT
from registry_sources import SOURCES


PERMIT_7 = ["Разрешение за строеж", "Влязло в сила на", "Издател", "Възложител", "Обект", "Район", "Местоположение"]
PERMIT_8 = ["Разрешение за строеж", "Поле 2", *PERMIT_7[1:]]
PERMIT_9 = ["Разрешение за строеж", "Поле 2", "Поле 3", *PERMIT_7[1:]]
CERTIFICATE_7 = ["Удостоверение", "Поле 2", "Издател", "Строителен надзор", "Обект", "Район", "Местоположение"]
VISA_13 = ["Поле 1", "Номер", "Дата", "Дата на влязъл в сила", "Основание", "Издател", "Район", "Обхват", "Виж на картата", "Поле 10", "Поле 11", "Вид документ", "Поле 13"]
VISA_14 = ["Поле 1", *VISA_13]

ALIASES = {
    "№": "number", "номер": "number", "удостоверение": "number", "виза": "number", "разрешение за строеж": "number",
    "дата": "date", "дата на издаване": "date",
    "влязло в сила на": "effective_date", "влязъл в сила": "effective_date", "влязъл в сила на": "effective_date",
    "дата на влязъл в сила": "effective_date", "дата на влизане в сила": "effective_date",
    "издател": "issuer", "възложител": "client", "надзор": "supervisor", "строителен надзор": "supervisor",
    "строеж/обект": "object", "строеж / обект": "object", "обект": "object",
    "район": "district", "обхват": "scope", "местоположение": "scope", "основание": "basis",
    "виж на картата": "map_link", "вид документ": "document_type", "виза за проектиране": "document_type",
}

DOCUMENT_CAPTIONS = {
    "разрешение за строеж",
    "удостоверение за въвеждане в експлоатация",
    "удостоверение за въвеждане в експлоатация на строеж",
    "виза за проектиране",
}


def cleaned(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def normalized_key(label: str) -> str:
    key = re.sub(r"\s+", " ", label.strip().casefold())
    return ALIASES.get(key, key)


def date_value(value, epoch=None):
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 25000 <= value <= 90000:
        converted = from_excel(value, epoch=epoch) if epoch is not None else from_excel(value)
        return converted.date().isoformat() if isinstance(converted, datetime) else None
    text = str(value or "")
    match = re.search(r"(?<!\d)(\d{1,2})\.(\d{1,2})\.(\d{4})(?!\d)", text)
    if match:
        try:
            return date(int(match[3]), int(match[2]), int(match[1])).isoformat()
        except ValueError:
            return None
    match = re.search(r"(?<!\d)(\d{4}-\d{2}-\d{2})(?!\d)", text)
    if match:
        try:
            return date.fromisoformat(match[1]).isoformat()
        except ValueError:
            return None
    return None


def identity(source_code: str, fields: dict) -> str:
    """Deduplicate equal values despite case/spacing and historical header names.

    No address matching, subset matching, or record-number-only replacement is
    used: any changed nonempty information remains a separate observation.
    """
    canonical = {}
    for label, value in fields.items():
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        key = normalized_key(label)
        text = re.sub(r"\s+", " ", str(value).strip()).casefold()
        if key == "number":
            lines = [re.sub(r"\s+", " ", line.strip()).casefold() for line in str(value).splitlines() if line.strip()]
            if lines:
                lines[0] = re.sub(r"^№\s*", "", lines[0])
            # The modern register embeds the document type and status under
            # its number. Only known type captions can be omitted. Statuses
            # such as "Влязъл в сила" are evidence and must distinguish rows.
            text = "\n".join(line for line in lines if line not in DOCUMENT_CAPTIONS)
        elif key in ("date", "effective_date"):
            text = date_value(value) or text
        # Repeated UI captions add no property evidence. Actual URLs and any
        # different document-type values remain part of the identity.
        if (key == "map_link" and text == "виж на картата") or (key == "document_type" and text == "виза за проектиране"):
            continue
        if key in canonical and canonical[key] != text:
            key = f"{key}:{label.casefold()}"
        canonical[key] = text
    payload = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256((source_code + "\0" + payload).encode("utf-8")).hexdigest()


def numeric_labels(code: str, headers: tuple) -> list[str]:
    indices = [int(value) for value in headers]
    if code == "building_permits":
        if indices == list(range(7)):
            return PERMIT_7
        if indices == list(range(8)):
            return PERMIT_8
        if all(0 <= index < 9 for index in indices):
            return [PERMIT_9[index] for index in indices]
    elif code == "exploitation_certificates" and all(0 <= index < 7 for index in indices):
        return [CERTIFICATE_7[index] for index in indices]
    elif code == "projecting_visas":
        if indices == list(range(13)):
            return VISA_13
        if indices == list(range(14)):
            return VISA_14
    raise ValueError(f"Unrecognized numeric column layout for {code}: {indices}")


def source_paths(root: Path):
    directories = (
        ("building_permits", "Razresheniya za Stroezh"),
        ("exploitation_certificates", "Udostoverenie za Vuvezhdane v Eksploataciya"),
        ("projecting_visas", "Vizi za Proektirane"),
    )
    for code, folder in directories:
        for path in sorted((root / folder).rglob("*.xlsx")):
            if path.name.startswith("~$"):
                continue
            actual_code = "exploitation_certificates_legacy" if "Old Registry" in path.parts else code
            yield actual_code, path


def workbook_records(code: str, path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, ())
        heavy = path.name == "Visas_of_Projection Full List Heavy Not Applicable.xlsx"
        if heavy:
            # Row one names parallel spreadsheet analysis blocks; row two is
            # the real header. Only columns 0–8 are original source records.
            # Filled/Blank QA counters (9–18) and manually merged helper cells
            # (19–27) have no proven official same-record provenance.
            header = next(iterator, ())
            if tuple(str(value).casefold() for value in header[:7]) != ("номер", "дата", "дата на влязъл в сила", "основание", "издател", "район", "обхват"):
                raise ValueError("Unrecognized analysis workbook layout")
            segments = [(0, 9)]
        else:
            segments = [(0, len(header))]
        if not heavy:
            is_numeric = all(isinstance(value, (int, float)) and not isinstance(value, bool) and int(value) == value for value in header)
            labels = numeric_labels(code, header) if is_numeric else [str(value).strip() if value is not None else f"Поле {index + 1}" for index, value in enumerate(header)]
        for row_number, row in enumerate(iterator, start=3 if heavy else 2):
            for start, end in segments:
                cells = row[start:end]
                segment_labels = [str(value).strip() if value is not None else f"Поле {index + 1}" for index, value in enumerate(header[start:end])] if heavy else labels
                fields = {label: cleaned(value) for label, value in zip(segment_labels, cells) if cleaned(value) is not None}
                if not fields:
                    continue
                yield row_number, start, fields, workbook.epoch
    finally:
        workbook.close()


def file_snapshot_date(path: Path):
    # Do not invent a year for names such as "Uploaded 21.01".
    match = re.search(r"(?<!\d)(\d{2})\.(\d{2})\.(\d{4})(?!\d)", path.name)
    return date_value(match[0]) if match else None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--previous", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True, help="Merged local version-one snapshot; use the sharded builder for publication.")
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--source-root", type=Path, default=ROOT)
    args = parser.parse_args()
    snapshot = read_json(args.previous, limit=MAX_LOCAL_BYTES)
    validate_snapshot(snapshot)
    identities = {identity(document["source_code"], document["raw_fields"]) for document in snapshot["documents"]}
    hashes = {(document["source_code"], document["document_hash"]) for document in snapshot["documents"]}
    sources = {source["code"] for source in snapshot["sources"]}
    files_seen = {}
    audit = {"previous_documents": len(snapshot["documents"]), "files": []}
    imported_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for code, path in source_paths(args.source_root):
        file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        result = {"source_code": code, "source_file": str(path.relative_to(args.source_root)), "sha256": file_hash, "observations": 0, "linked_observations": 0, "duplicates": 0, "added": 0}
        if (code, file_hash) in files_seen:
            result["identical_file"] = files_seen[(code, file_hash)]
            audit["files"].append(result)
            print(json.dumps(result, ensure_ascii=False), flush=True)
            continue
        files_seen[(code, file_hash)] = result["source_file"]
        if code not in sources:
            source = SOURCES[code]
            snapshot["sources"].append({key: source[key] for key in ("label_bg", "label_en", "source_url")} | {"code": code, "snapshot_file": None, "snapshot_date": None})
            sources.add(code)
        snapshot_date = file_snapshot_date(path)
        for row_number, segment, fields, epoch in workbook_records(code, path):
            result["observations"] += 1
            serialized = json.dumps(fields, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            ids = sorted(set(CADASTRAL_ID.findall(serialized)))
            if not ids:
                continue
            result["linked_observations"] += 1
            fingerprint = identity(code, fields)
            digest = hashlib.sha256((code + "\0" + serialized).encode("utf-8")).hexdigest()
            if fingerprint in identities or (code, digest) in hashes:
                result["duplicates"] += 1
                continue
            number = next((value for label, value in fields.items() if normalized_key(label) == "number"), None)
            issued_at = date_value(number)
            if issued_at is None:
                issued_at = next((date_value(value, epoch) for label, value in fields.items() if normalized_key(label) == "date" and date_value(value, epoch) is not None), None)
            snapshot["documents"].append({"source_code": code, "document_hash": digest, "record_key": str(number)[:190] if number is not None else None, "raw_fields": fields, "snapshot_file": path.name, "snapshot_date": snapshot_date, "issued_at": issued_at, "cadastral_ids": ids})
            identities.add(fingerprint)
            hashes.add((code, digest))
            result["added"] += 1
        snapshot["imports"].append({"source_code": code, "source_file": path.name, "snapshot_date": snapshot_date, "imported_at": imported_at, "records_loaded": result["added"]})
        audit["files"].append(result)
        print(json.dumps(result, ensure_ascii=False), flush=True)
    snapshot["generated_at"] = imported_at
    payload = json_bytes(snapshot)
    ids = {plot_id for document in snapshot["documents"] for plot_id in document["cadastral_ids"]}
    audit.update({"documents": len(snapshot["documents"]), "added_documents": len(snapshot["documents"]) - audit["previous_documents"], "json_bytes": len(payload), "cadastral_identifiers": len(ids), "parcel_document_references": sum(len(document["cadastral_ids"]) for document in snapshot["documents"])})
    write_atomic(args.audit, json_bytes(audit))
    print(json.dumps({key: value for key, value in audit.items() if key != "files"}, ensure_ascii=False), flush=True)
    # Report the full archive result before any bounds failure. Publication is
    # never made with a silently truncated subset of the available evidence.
    validate_snapshot(snapshot)
    if len(payload) > MAX_LOCAL_BYTES:
        raise ValueError(f"Merged snapshot exceeds the {MAX_LOCAL_BYTES:,}-byte local processing bound; full audit saved.")
    write_atomic(args.output, payload)


if __name__ == "__main__":
    main()
