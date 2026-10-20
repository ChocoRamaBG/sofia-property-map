#!/usr/bin/env python3
"""Build an offline grouped street/quarter reference, without parcel IDs.

Read the 24 verified cached KAIS land-parcel exports and an existing coordinate
workbook. Exact cadastral IDs are used only inside this process. Public output
contains aggregated reference circles, not properties, boundaries or owners.
No network requests are performed. The output is reusable by daily builds.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import statistics
import struct
from zipfile import ZipFile

from imot_location import MAX_REFERENCE_BYTES, SOURCE_URL, distance_m, neighborhood_name, stable_key, street_name
from refresh_cadastre import CADASTRAL_ID, archive_records


def file_digest(path):
    result = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(128 * 1024), b""):
            result.update(block)
    return result.hexdigest()


def structured_rows(path):
    """Read only safe structured address fields; archive_records validates DBF."""
    with ZipFile(path) as archive:
        names = [n for n in archive.namelist() if n.lower().endswith(".dbf")]
        if len(names) != 1:
            raise ValueError("One cadastral DBF is required")
        with archive.open(names[0]) as handle:
            header = handle.read(32)
            count = struct.unpack("<I", header[4:8])[0]
            header_length, record_length = struct.unpack("<HH", header[8:12])
            fields = []
            for _ in range((header_length - 33) // 32):
                field = handle.read(32)
                fields.append((field[:11].split(b"\0")[0].decode("ascii"), field[16]))
            handle.read(1)
            for _ in range(count):
                raw, position, values = handle.read(record_length), 1, {}
                for name, length in fields:
                    if name in {"cadnum", "quarname", "strename"}:
                        values[name] = raw[position:position + length].decode("utf-8").strip(" \x00")
                    position += length
                yield values


def coordinate_points(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    if tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))) != ("ID", "X", "Y"):
        raise ValueError("Coordinate workbook headings changed")
    points, conflicting, rejected = {}, set(), 0
    for identifier, lat, lng in sheet.iter_rows(min_row=2, values_only=True):
        try:
            if not isinstance(identifier, str) or not CADASTRAL_ID.fullmatch(identifier):
                raise ValueError("Invalid point ID")
            lat, lng = float(lat), float(lng)
            if not math.isfinite(lat) or not math.isfinite(lng) or not (42.4 < lat < 43 and 22.8 < lng < 24):
                raise ValueError("Point outside reference extent")
            point = (lat, lng)
            if identifier in points and points[identifier] != point:
                conflicting.add(identifier)
            points[identifier] = point
        except (ValueError, TypeError):
            rejected += 1
    workbook.close()
    for identifier in conflicting:
        points.pop(identifier, None)
    if not 100000 <= len(points) <= 500000:
        raise ValueError("Coordinate reference count requires inspection")
    return points, {"coordinate_point_count": len(points), "rejected_coordinate_rows": rejected, "conflicting_coordinate_ids": len(conflicting)}


def circle(points, precision):
    unique = sorted(set(points))
    if len(unique) < 3:
        return None
    lat, lng = statistics.median(x[0] for x in unique), statistics.median(x[1] for x in unique)
    excluded = 0
    if precision == "neighborhood" and len(unique) >= 20:
        distances = sorted(distance_m(lat, lng, x, y) for x, y in unique)
        cutoff = max(1800, 3 * distances[int((len(distances) - 1) * .9)])
        retained = [p for p in unique if distance_m(lat, lng, *p) <= cutoff]
        # A small set of distant source points can make a coherent quarter
        # unusable. Retain only a >=95% dominant central group, recording every
        # excluded point. Multi-area labels are not reduced to a guessed area.
        if len(retained) >= math.ceil(.95 * len(unique)):
            excluded = len(unique) - len(retained)
            unique = retained
            lat, lng = statistics.median(x[0] for x in unique), statistics.median(x[1] for x in unique)
    # Coordinates are deliberately rounded (~100 m), and the radius includes
    # every input point plus a margin. It is an uncertainty circle, not a boundary.
    lat, lng = round(lat, 3), round(lng, 3)
    maximum = max(distance_m(lat, lng, x, y) for x, y in unique)
    radius = math.ceil((maximum + (400 if precision == "neighborhood" else 150)) / 100) * 100
    radius = max((1000 if len(unique) < 10 else 700) if precision == "neighborhood" else 250, radius)
    if radius > (8000 if precision == "neighborhood" else 6000):
        return None
    return {"lat": lat, "lng": lng, "radius_m": radius, "point_count": len(unique), "excluded_outlier_point_count": excluded}


def build(args):
    manifest = json.loads(args.cadastre_manifest.read_text(encoding="utf-8"))
    if manifest.get("version") != 2 or manifest.get("completion", {}).get("district_count") != 24 or len(manifest.get("sources", [])) != 24:
        raise ValueError("A complete 24-district cadastral manifest is required")
    coordinates, audit = coordinate_points(args.coordinates)
    selected, conflicts, source_meta = {}, set(), []
    raw_quarters, matched_streets = Counter(), Counter()
    source_total = 0
    for index, source in enumerate(manifest["sources"]):
        path = args.archives / f"sofia-district-{index:02d}.zip"
        if path.is_symlink() or path.stat().st_size != source["zip_bytes"] or file_digest(path) != source["zip_sha256"]:
            raise ValueError("Cached cadastral archive does not match its verified manifest")
        district = source["label_bg"].split(" - ", 1)[1]
        seen = 0
        safe_rows = archive_records(path, district)
        for validated, structured in zip(safe_rows, structured_rows(path), strict=True):
            identifier, _, safe = validated
            if identifier != structured["cadnum"]:
                raise ValueError("Structured address and verified row do not agree")
            seen += 1
            raw_quarters[structured["quarname"]] += 1
            if identifier not in coordinates:
                continue
            item = (coordinates[identifier], neighborhood_name(structured["quarname"]), street_name(structured["strename"]), district)
            if identifier in selected and selected[identifier] != item:
                conflicts.add(identifier)
            else:
                selected[identifier] = item
        if seen != source["record_count"]:
            raise ValueError("Archive row count does not reconcile with manifest")
        source_total += seen
        source_meta.append({k: source[k] for k in ("source_url", "zip_sha256", "captured_at", "catalogue_published_at") if source.get(k)})
    conflicts.update(manifest.get("ambiguous_ids", []))
    for identifier in conflicts:
        selected.pop(identifier, None)
    places, streets = defaultdict(list), defaultdict(list)
    place_districts, street_places = defaultdict(set), defaultdict(set)
    for point, quarter, street, district in selected.values():
        if quarter:
            places[quarter].append(point)
            place_districts[quarter].add(district)
        if street:
            streets[(street, district)].append(point)
            matched_streets[street] += 1
            if quarter:
                street_places[(street, district)].add(quarter)
    public_places, public_streets, valid_places = [], [], {}
    for name in sorted(places):
        geometry = circle(places[name], "neighborhood")
        if geometry:
            key = stable_key("place", name)
            valid_places[name] = key
            public_places.append({"key": key, "label_bg": name.title(), **geometry, "districts_bg": sorted(place_districts[name])})
    prefixes = {"street":"ул.", "boulevard":"бул.", "square":"пл."}
    for (street, district), points in sorted(streets.items()):
        geometry = circle(points, "street")
        if geometry:
            kind, name = street
            key = stable_key("street", kind + ":" + name + ":" + district.casefold())
            public_streets.append({"key": key, "label_bg": prefixes[kind] + " " + name.title(), **geometry,
                                  "district_bg": district, "place_keys": sorted(valid_places[x] for x in street_places[(street, district)] if x in valid_places)})
    audit.update({"official_source_rows": source_total, "raw_quarter_labels": len(raw_quarters), "rows_without_quarter": raw_quarters[""],
                  "exact_id_matched_points": len(selected), "ambiguous_source_ids_excluded": len(conflicts),
                  "grouped_neighborhood_count": len(public_places), "grouped_street_count": len(public_streets),
                  "street_reference_point_count": sum(v for v in matched_streets.values())})
    result = {"version": 1, "reference_date": manifest["generated_at"][:10],
              "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z"),
              "source_url": SOURCE_URL, "attribution_bg": "АГКК, КАИС Отворени данни — групирани адресни ориентири. Приблизителни локации, не граници или идентификация на имот.",
              "method": "Exact official identifiers are used internally to associate existing coordinate points with public address fields, then discarded. Median grouped centers rounded to 0.001 degrees; circles enclose retained reference points plus a margin, not certified quarter boundaries. For quarters with at least 20 points, distant outliers can be excluded only when at least 95 percent remain in the central group, with excluded counts published. Groups need at least 3 distinct points. No advert is joined to an official parcel. Missing, sparse or ambiguous reference groups are unresolved.",
              "coverage": audit, "sources": source_meta, "places": public_places, "streets": public_streets}
    payload = (json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")
    if not 1 <= len(payload) <= MAX_REFERENCE_BYTES:
        raise ValueError("Public market reference exceeds its 4 MiB limit")
    # Serialize only allowlisted aggregates; raw IDs, addresses, street numbers,
    # owners, contact fields and workbook names never enter the output object.
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(payload)
    return {**audit, "reference_bytes": len(payload)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archives", required=True, type=Path)
    parser.add_argument("--cadastre-manifest", required=True, type=Path)
    parser.add_argument("--coordinates", required=True, type=Path)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "public/data/market-georeference.json")
    args = parser.parse_args()
    print(json.dumps(build(args), ensure_ascii=False))


if __name__ == "__main__":
    main()
