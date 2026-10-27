#!/usr/bin/env python3
"""Prepare transparent asking-price analytics from complete imot.bg captures.

Historical Excel import is read-only. Its September 2024 dates are retained;
generation time never turns old rows into current offers or daily alerts.
"""
from __future__ import annotations

import argparse
import gzip
import io
from collections import Counter, defaultdict
from datetime import datetime
import json
from pathlib import Path
import re
import statistics
from uuid import uuid4

from imot_location import LocationResolver

from imot_market_format import (
    VERSION, MAX_LISTINGS, MAX_SHARD_BYTES, MAX_SUMMARY_BYTES, MAX_MANIFEST_BYTES,
    MAX_STATE_BYTES, MAX_STATE_COMPRESSED_BYTES, MAX_CAPTURE_BYTES, MAX_SHARDS, SHARD_ROWS, canonical_offer_id, digest,
    json_bytes, number, stable_key, text, timestamp, utc_now, validate_snapshot,
)

MIN_PEERS = 8
DEAL_DISCOUNT_PCT = 20
EXTREME_DISCOUNT_PCT = 60
COMPARISON_METHOD_VERSION = 3
AREA_TOLERANCE = 0.25
YEAR_TOLERANCE = 10
FLOOR_TOLERANCE = 2
BGN_PER_EUR = 1.95583
VAT_LABELS = {"included": "Включен ДДС / не се начислява", "excluded": "Без ДДС", "unknown": "ДДС не е посочен"}


def normalized_analysis(raw, label, checked_at):
    source = raw.get("analysis") if isinstance(raw.get("analysis"), dict) else {}
    stage = source.get("completion_stage", "unknown")
    stated_stage = "unknown"
    label = (label or "").casefold()
    if re.search(r"ще\s+бъде|планиран|в строеж|в процес|на проект|акт\s*1[45]", label):
        stated_stage = "planned"
    elif re.search(r"въведен[ао]?\s+в\s+експлоатация|завършен", label):
        stated_stage = "completed"
    conflict = stage in {"planned", "completed"} and stated_stage in {"planned", "completed"} and stage != stated_stage
    stage = "unknown" if conflict else stated_stage if stated_stage != "unknown" else stage
    allowed = {
        "completion_stage": {"planned", "completed", "unknown"},
        "condition": {"renovated", "needs_renovation", "unknown"},
        "area_basis": {"net", "gross", "terrace_included", "unknown"},
    }
    answer = {field: source.get(field) if source.get(field) in values else "unknown" for field, values in allowed.items()}
    # Earlier captures recognized area words without establishing which number
    # they described. A net-area mention must not relabel a gross-area quote.
    if raw.get("analysis_method_version") not in {2, 3}:
        answer["area_basis"] = "unknown"
    answer["completion_stage"] = stage if stage in allowed["completion_stage"] else "unknown"
    risks = {"share_sale", "auction", "tenanted", "nonresidential", "mandatory_extra_cost"}
    answer["risk_flags"] = sorted({x for x in source.get("risk_flags", []) if isinstance(x, str) and x in risks})
    answer["checked_at"] = timestamp(source.get("checked_at")) or checked_at
    answer["description_analyzed"] = source.get("description_analyzed") is True
    return answer, conflict


def floor_band(value):
    raw = (value or "").casefold()
    if re.search(r"сутерен|партер|приземен|подзем", raw):
        return "ground_or_below"
    match = re.search(r"(?<!\d)(\d{1,2})(?:\s*-?\s*(?:ви|ри|ти|ми))?\s*(?:ет\.?\s*)?(?:от|/)\s*(\d{1,2})(?!\d)", raw)
    if match:
        current, total = map(int, match.groups())
        if 0 <= current <= total <= 60:
            return "ground_or_below" if current == 0 else "top" if current == total else "middle"
    return "unknown"


def normalized_features(raw, observed):
    source = raw.get("property_features") if isinstance(raw.get("property_features"), dict) else {}
    allowed = {
        "construction_type": {"brick", "panel", "epk", "monolithic", "other", "unknown"},
        "furnishing": {"furnished", "partial", "unfurnished", "unknown"},
        "year_basis": {"built", "commissioned", "planned", "unknown"},
        "elevator": {"yes", "no", "unknown"},
    }
    answer = {field: source.get(field) if source.get(field) in values else "unknown" for field, values in allowed.items()}
    for field, minimum, maximum in (("building_year", 1800, 2100), ("floor_number", -5, 100), ("total_floors", 1, 100)):
        value = source.get(field)
        answer[field] = value if type(value) is int and minimum <= value <= maximum else None
    if answer["floor_number"] is not None and answer["total_floors"] is not None and answer["floor_number"] > answer["total_floors"]:
        answer["floor_number"] = answer["total_floors"] = None
    if answer["building_year"] is None:
        answer["year_basis"] = "unknown"
    answer["checked_at"] = timestamp(source.get("checked_at")) or observed
    answer["source"] = source.get("source") if source.get("source") in {"card", "description"} else "card"
    return answer


def normalized_agency(raw, seller):
    source = raw.get("agency")
    if seller != "agency" or not isinstance(source, dict) or source.get("evidence") != "professional_profile":
        return None
    # The closed validator checks URLs/contact form before publication.
    answer = {field: source.get(field) for field in ("name", "profile_url", "phone", "website", "checked_at", "evidence")}
    answer["name"] = text(answer["name"], 120)
    answer["checked_at"] = timestamp(answer["checked_at"])
    return answer if answer["name"] and answer["profile_url"] else None


def offer_floor_band(offer):
    features = offer["property_features"]
    current, total = features["floor_number"], features["total_floors"]
    if current is not None and current <= 0:
        return "ground_or_below"
    if current is not None and total is not None:
        return "top" if current == total else "middle"
    return floor_band(offer["floor_bg"])


def actual_building_year(offer):
    features = offer["property_features"]
    return features["building_year"] if features["year_basis"] in {"built", "commissioned"} else None


def normalize_vat(value):
    raw = (text(value) or "").casefold()
    if raw in {"included", "not_applicable", "included_or_not_applicable"} or raw == "с ддс" or "включено ддс" in raw or "не се начислява" in raw:
        return "included"
    if raw == "excluded" or "без ддс" in raw:
        return "excluded"
    return "unknown"


def normalize_currency(value):
    raw = (text(value) or "").upper()
    if raw in {"EUR", "EURO", "€", "ЕВРО"}:
        return "EUR"
    if raw in {"BGN", "ЛВ.", "ЛВ", "ЛЕВА"}:
        return "BGN"
    return None


def normalize_neighborhood(value):
    raw = text(value)
    if not raw:
        return None
    raw = re.sub(r"^(?:гр\.?\s*София\s*[,\-]?\s*)", "", raw, flags=re.I)
    raw = re.sub(r"^(?:кв\.?|ж\.?\s*к\.?)\s*", "", raw, flags=re.I).strip()
    raw = " ".join(raw.split())
    if raw.isupper():
        raw = raw.title()
    return raw or None


def normalize_offer(raw, default_observed, historical, resolver=None):
    identifier, url = canonical_offer_id(raw.get("url"), raw.get("id"))
    neighborhood = normalize_neighborhood(raw.get("neighborhood", raw.get("neighborhood_bg")))
    kind = text(raw.get("property_type", raw.get("type_bg")), 64)
    if kind:
        kind = kind.upper().replace("–", "-").replace("- ", "-")
    area = number(raw.get("area_m2", raw.get("area_sqm")))
    original_price = number(raw.get("price"))
    currency = normalize_currency(raw.get("currency", raw.get("price_currency")))
    price = original_price / BGN_PER_EUR if original_price and currency == "BGN" else original_price if currency == "EUR" else None
    eur_sqm = price / area if price and area else None
    derived = raw.get("price_total_derived", False)
    if not isinstance(derived, bool):
        raise ValueError("Derived total marker must be boolean")
    unit_quote = number(raw.get("quoted_eur_per_sqm")) if derived else None
    if derived:
        if unit_quote is None or area is None or currency not in {"EUR", "BGN"}:
            raise ValueError("Derived total requires an explicit EUR unit quote and area")
        # The source quoted a rate; this multiplication is our calculated total.
        price, eur_sqm = unit_quote * area, unit_quote
    year = number(raw.get("construction_year"))
    year = int(year) if year and year.is_integer() and 1800 <= year <= 2100 else None
    construction = "new_build_2021_plus" if year and year >= 2021 else "pre_2021" if year else "unknown"
    construction_label = text(raw.get("construction_label_bg"), 80)
    if year is None and construction_label:
        # Published decade bands can identify a cohort without inventing a year.
        band = re.search(r"\b((?:18|19|20)\d{2})\s*[-–—]\s*((?:18|19|20)\d{2})\b", construction_label)
        if band:
            low, high = map(int, band.groups())
            if low <= high < 2021:
                construction = "pre_2021"
            elif 2021 <= low <= high:
                construction = "new_build_2021_plus"
    quality = []
    if not neighborhood:
        quality.append("missing_neighborhood")
    if not kind:
        quality.append("missing_type")
    if not currency:
        quality.append("missing_currency")
    if not original_price and not (derived and unit_quote):
        quality.append("missing_price")
    elif not currency:
        quality.append("unconverted_price_currency")
    if not area:
        quality.append("missing_area")
    if (area and not 10 <= area <= 2000) or (price and not 1000 <= price <= 20000000) or (eur_sqm and not 50 <= eur_sqm <= 50000):
        quality.append("implausible_units")
    observed = timestamp(raw.get("observed_at")) or (None if historical else default_observed)
    if historical and observed is None:
        quality.append("unknown_observed_at")
    seller = raw.get("seller_type")
    evidence = raw.get("seller_evidence")
    if (seller, evidence) not in {("agency", "agency_profile"), ("private", "explicit_private_claim"), ("unknown", "not_reported"), ("unknown", "conflicting_signals")}:
        seller, evidence = "unknown", "not_reported"
    analysis, conflicting_completion = normalized_analysis(raw, construction_label, observed)
    if conflicting_completion:
        quality.append("conflicting_completion_claims")
    location = resolver.resolve(neighborhood_bg=neighborhood, address_bg=raw.get("address_bg"), location_hints=raw.get("location_hints")) if resolver else None
    return {
        "id": identifier, "url": url,
        "neighborhood_key": stable_key("neighborhood", neighborhood) if neighborhood else None,
        "neighborhood_bg": neighborhood,
        "type_key": stable_key("type", kind) if kind else None, "type_bg": kind,
        "area_sqm": round(area, 4) if area else None,
        "price_eur": round(price, 2) if price else None,
        "eur_per_sqm": round(eur_sqm, 4) if eur_sqm else None,
        "price_total_derived": derived,
        "quoted_eur_per_sqm": round(unit_quote, 4) if unit_quote else None,
        "price_currency": currency,
        "vat_basis": normalize_vat(raw.get("vat", raw.get("vat_basis"))),
        "construction_year": year, "construction_cohort": construction,
        "floor_bg": text(raw.get("floor_bg"), 80),
        "construction_label_bg": construction_label,
        "address_bg": text(raw.get("address_bg"), 160),
        "last_modified_at": timestamp(raw.get("updated_at", raw.get("last_modified_at"))),
        "observed_at": observed, "first_seen_at": observed,
        "comparison": None, "change": None, "historical": historical,
        "quality_flags": quality,
        "seller_type": seller, "seller_evidence": evidence,
        "seller_checked_at": timestamp(raw.get("seller_checked_at")),
        "agency": normalized_agency(raw, seller),
        "property_features": normalized_features(raw, observed),
        "analysis": analysis, "location": location, "asking_estimate": None,
    }


def date_from_legacy(value):
    if not isinstance(value, str):
        return None
    try:
        return datetime.strptime(value.strip(), "%d-%m-%y_%H:%M").date().isoformat()
    except ValueError:
        return None


def workbook_capture(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    master_dates = {}
    for row in workbook["Master"].iter_rows(min_row=3, values_only=True):
        if row[0] and len(row) > 12 and date_from_legacy(row[12]):
            master_dates[str(row[0]).casefold()] = date_from_legacy(row[12])
    offers, sources, observed_dates = [], [], []
    for sheet in workbook:
        if sheet.title in {"Master", "Eligible Offers"}:
            continue
        header = next(sheet.values)
        if len(header) < 16 or header[0] != "Квартал" or header[15] != "Линк към обявата":
            raise ValueError("Historical worksheet has an unexpected schema")
        source_date = master_dates.get(sheet.title.casefold())
        if sheet.title.startswith("Malinova "):
            source_date = master_dates.get("malinova dolina")
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        for row in rows:
            for value in row[16:22]:
                detected = date_from_legacy(value)
                if detected:
                    source_date = detected
                    break
        count = 0
        for row in rows:
            if not isinstance(row[15], str) or "imot.bg" not in row[15]:
                continue
            offers.append({
                "url": row[15], "neighborhood": row[0], "property_type": row[1],
                "construction_year": row[2], "price": row[3], "vat": row[4],
                "currency": row[5], "area_m2": row[6], "updated_at": row[11],
                "observed_at": source_date,
            })
            count += 1
        sources.append({"key": stable_key("legacy-sheet", sheet.title), "captured_count": count, "observed_at": source_date})
        if source_date:
            observed_dates.append(source_date)
    workbook.close()
    return {
        "version": VERSION, "source": "imot.bg", "complete": False,
        "mode": "historical_bootstrap", "started_at": min(observed_dates), "completed_at": max(observed_dates),
        "scope": {"city": "Sofia", "transaction": "sale", "property_types": ["2-СТАЕН", "3-СТАЕН", "4-СТАЕН"], "construction_year_min": 2021},
        "completion": {"sources": sources, "total_rows": len(offers)}, "offers": offers,
        "audit": {"source_file_sha256": digest(Path(path).read_bytes()), "source_sheet_count": len(sources)},
    }


def stats(values):
    if not values:
        return {"count": 0, "median_eur_per_sqm": None, "mean_eur_per_sqm": None, "trimmed_mean_eur_per_sqm": None}
    values = sorted(values)
    trim = int(len(values) * 0.1) if len(values) >= 20 else 0
    central = values[trim:len(values) - trim] if trim else values
    return {"count": len(values), "median_eur_per_sqm": round(statistics.median(values), 2), "mean_eur_per_sqm": round(statistics.mean(values), 2), "trimmed_mean_eur_per_sqm": round(statistics.mean(central), 2)}


def comparable(offer):
    return not offer["quality_flags"] and offer["type_bg"] in {"2-СТАЕН", "3-СТАЕН", "4-СТАЕН"}


def vat_statistics(offers):
    result = []
    for basis in ("included", "excluded", "unknown"):
        rows = [x for x in offers if x["vat_basis"] == basis]
        metrics = stats([x["eur_per_sqm"] for x in rows])
        metrics.pop("count", None)
        result.append({"vat_basis": basis, "label_bg": VAT_LABELS[basis], "priced_count": len(rows), **metrics, "median_price_eur": round(statistics.median([x["price_eur"] for x in rows]), 2) if rows else None})
    return result


def cohort_key(offer):
    features = offer["property_features"]
    return ":".join([offer["neighborhood_key"] or "unknown", offer["type_key"] or "unknown", offer["vat_basis"], offer["construction_cohort"], offer["analysis"]["completion_stage"], offer["analysis"]["area_basis"], offer_floor_band(offer), features["construction_type"], features["furnishing"], features["elevator"], "actual_year" if actual_building_year(offer) is not None else "unverified_year"])


def peer_sample(offer, peers):
    """Property text supplies claims, not independently verified attributes."""
    rows, seen, duplicates = [], {offer["id"]}, 0
    year = actual_building_year(offer)
    floor = offer["property_features"]["floor_number"]
    for row in peers:
        if row["id"] in seen:
            if row["id"] != offer["id"]:
                duplicates += 1
            continue
        seen.add(row["id"])
        peer_year = actual_building_year(row)
        peer_floor = row["property_features"]["floor_number"]
        if (not offer["area_sqm"] * (1 - AREA_TOLERANCE) <= row["area_sqm"] <= offer["area_sqm"] * (1 + AREA_TOLERANCE)
                or (offer["analysis"]["condition"] != "unknown" and row["analysis"]["condition"] != offer["analysis"]["condition"])
                or (year is not None and (peer_year is None or abs(year - peer_year) > YEAR_TOLERANCE))
                or (floor is not None and (peer_floor is None or abs(floor - peer_floor) > FLOOR_TOLERANCE))):
            continue
        rows.append(row)
    location = offer.get("location") or {}
    if location.get("precision") == "street":
        local = [x for x in rows if (x.get("location") or {}).get("key") == location["key"]]
        if len(local) >= MIN_PEERS:
            return local, "street", duplicates
    return rows, "neighborhood", duplicates


def estimate_evidence(offer, location_basis, values, duplicates):
    known = {
        "neighborhood": bool(offer["neighborhood_key"]), "type": bool(offer["type_key"]),
        "vat_basis": offer["vat_basis"] != "unknown", "construction_age": actual_building_year(offer) is not None,
        "construction_type": offer["property_features"]["construction_type"] != "unknown",
        "completion_stage": offer["analysis"]["completion_stage"] != "unknown",
        "area_basis": offer["analysis"]["area_basis"] != "unknown", "floor": offer["property_features"]["floor_number"] is not None and offer["property_features"]["total_floors"] is not None,
        "furnishing": offer["property_features"]["furnishing"] != "unknown",
        "condition": offer["analysis"]["condition"] != "unknown", "elevator": offer["property_features"]["elevator"] != "unknown",
        "area_tolerance": bool(offer["area_sqm"]), "street": location_basis == "street",
    }
    missing = [field for field, available in known.items() if not available]
    detailed_fields = [field for field in known if field != "street"]
    supported = len(values) >= MIN_PEERS
    median = statistics.median(values) if values else None
    low, high = (quartile(values, 0.25), quartile(values, 0.75)) if supported else (None, None)
    position = "unavailable" if not supported else "below_range" if offer["eur_per_sqm"] < low else "above_range" if offer["eur_per_sqm"] > high else "within_range"
    return {"quality": "insufficient" if not supported else "detailed" if all(known[field] for field in detailed_fields) else "limited",
            "matched_fields": [field for field, available in known.items() if available], "missing_fields": missing,
            "spread_pct": round((high - low) / median * 100, 2) if supported and median else None,
            "price_position": position, "duplicate_peer_count": duplicates, "unknown_physical_duplicates": True}


def quartile(values, fraction):
    ordered = sorted(values)
    point = (len(ordered) - 1) * fraction
    low = int(point)
    high = min(low + 1, len(ordered) - 1)
    return ordered[low] + (ordered[high] - ordered[low]) * (point - low)


def build_snapshot(capture, output, previous=None):
    if isinstance(capture.get("capture"), dict) and isinstance(capture.get("listings"), list):
        context = capture["capture"]
        capture = {
            "version": capture.get("version"), "source": context.get("source"),
            "mode": context.get("mode", "daily"), "started_at": context.get("started_at"),
            "completed_at": context.get("finished_at"), "complete": context.get("complete"),
            "scope": context.get("scope", {}),
            "scope_bg": context.get("scope_bg"), "source_url": context.get("source_url"),
            "requested_neighborhood_count": context.get("requested_neighborhood_count"),
            "captured_neighborhood_count": context.get("captured_neighborhood_count"),
            "completion": {"sources": context.get("searches", []), "total_unique": context.get("captured_count")},
            "offers": capture["listings"], "audit": {"capture": context},
            "enrichment_cache": capture.get("enrichment_cache", {}),
            "enrichment": capture.get("enrichment", {}),
        }
    historical = capture.get("mode") == "historical_bootstrap"
    if type(capture.get("version")) is not int or capture["version"] != VERSION or capture.get("source") != "imot.bg" or not isinstance(capture.get("offers"), list) or len(capture["offers"]) > MAX_LISTINGS * 2:
        raise ValueError("Unsupported or excessive source capture")
    captured_at = timestamp(capture.get("completed_at"))
    if captured_at is None:
        raise ValueError("Source capture lacks a valid completion date")
    if not historical:
        completion = capture.get("completion", {})
        sources = completion.get("sources")
        if capture.get("complete") is not True or not isinstance(sources, list) or not sources:
            raise ValueError("Only a complete daily source capture can be published")
        if not timestamp(capture.get("started_at")) or capture["started_at"] > capture["completed_at"]:
            raise ValueError("Source capture chronology is invalid")
        source_keys = set()
        for source in sources:
            if type(source.get("expected_count")) is not int or type(source.get("captured_count")) is not int or source.get("expected_count") != source.get("captured_count") or source["expected_count"] < 0:
                raise ValueError("A daily source has not reconciled its declared offer count")
            key = source.get("key")
            if not isinstance(key, str) or not key or key in source_keys:
                raise ValueError("Daily source keys are missing or repeated")
            source_keys.add(key)
    resolver = LocationResolver()
    normalized = []
    rejected = 0
    for raw in capture["offers"]:
        try:
            normalized.append(normalize_offer(raw, captured_at, historical, resolver))
        except (ValueError, TypeError, AttributeError):
            rejected += 1
    if not historical and rejected:
        raise ValueError("Daily capture contains offers without valid stable source IDs")
    selected, versions = {}, defaultdict(list)
    for offer in normalized:
        versions[offer["id"]].append(offer)
        existing = selected.get(offer["id"])
        if existing is None or (offer["observed_at"] or "", offer["last_modified_at"] or "", json_bytes(offer)) > (existing["observed_at"] or "", existing["last_modified_at"] or "", json_bytes(existing)):
            selected[offer["id"]] = offer
    for identifier, variants in versions.items():
        latest = selected[identifier]
        latest_values = {(x["neighborhood_key"], x["type_key"], x["area_sqm"], x["price_eur"], x["price_currency"], x["vat_basis"], x["price_total_derived"], x["quoted_eur_per_sqm"]) for x in variants if x["observed_at"] == latest["observed_at"]}
        if len(latest_values) > 1:
            latest["quality_flags"].append("duplicate_conflict")
    listings = sorted(selected.values(), key=lambda x: x["id"])
    if not 1 <= len(listings) <= MAX_LISTINGS:
        raise ValueError("Market dataset is empty or exceeds the listing bound")
    if not historical and capture["completion"].get("total_unique") != len(listings):
        raise ValueError("Complete daily unique listing total does not reconcile")
    peers = defaultdict(list)
    for offer in listings:
        if comparable(offer) and not offer["analysis"]["risk_flags"]:
            peers[cohort_key(offer)].append(offer)
    for offer in listings:
        key = cohort_key(offer)
        rows, location_basis, duplicates = peer_sample(offer, peers.get(key, [])) if comparable(offer) else ([], "neighborhood", 0)
        values = [x["eur_per_sqm"] for x in rows]
        comparison = {"method_version": COMPARISON_METHOD_VERSION, "cohort_key": key, "peer_count": len(values), "eligible": False, "alert_eligible": False, "vat_basis_known": offer["vat_basis"] != "unknown", "construction_year_known": offer["construction_year"] is not None, "construction_cohort_known": offer["construction_cohort"] != "unknown", "reason": "invalid_comparable_fields"}
        estimate = {"status": "invalid_inputs", "method_version": COMPARISON_METHOD_VERSION, "peer_count": len(values), "location_basis": location_basis,
                    "median_eur_per_sqm": None, "p25_eur_per_sqm": None, "p75_eur_per_sqm": None,
                    "indicated_price_eur": None, "range_low_eur": None, "range_high_eur": None,
                    "evidence": estimate_evidence(offer, location_basis, values, duplicates)}
        if comparable(offer):
            metrics = stats(values)
            median = statistics.median(values) if values else None
            discount = round((1 - offer["eur_per_sqm"] / median) * 100, 2) if median else None
            comparison.update({**metrics, "peer_count": len(values), "below_median_pct": discount})
            comparison.pop("count", None)
            if offer["analysis"]["risk_flags"]:
                comparison["reason"] = "advertised_terms_review"
                estimate["status"] = "needs_review"
            elif len(values) < MIN_PEERS:
                comparison["reason"] = "insufficient_peers"
                estimate["status"] = "insufficient_comparables"
            elif discount is not None and discount >= EXTREME_DISCOUNT_PCT:
                comparison["reason"] = "extreme_discount_review"
                estimate["status"] = "needs_review"
            elif discount is not None and discount >= DEAL_DISCOUNT_PCT:
                comparison.update({"eligible": True, "reason": "below_peer_asking_median"})
                estimate["status"] = "available"
            else:
                comparison["reason"] = "below_threshold"
                estimate["status"] = "available"
            if len(values) >= MIN_PEERS and median:
                low, high = quartile(values, 0.25), quartile(values, 0.75)
                estimate.update({"median_eur_per_sqm": round(median, 2), "p25_eur_per_sqm": round(low, 2), "p75_eur_per_sqm": round(high, 2),
                                 "indicated_price_eur": round(median * offer["area_sqm"], 2), "range_low_eur": round(low * offer["area_sqm"], 2), "range_high_eur": round(high * offer["area_sqm"], 2)})
        comparison["alert_eligible"] = comparison["eligible"] and comparison["vat_basis_known"]
        if estimate["status"] != "available":
            estimate["evidence"]["price_position"] = "unavailable"
        offer["comparison"] = comparison
        offer["asking_estimate"] = estimate
    previous = previous or {}
    if previous and previous.get("version") != VERSION:
        raise ValueError("Unsupported previous market state")
    prior = previous.get("listings", {})
    initial_baseline = historical or previous.get("mode") != "daily"
    methodology_changed = previous.get("comparison_method_version") != COMPARISON_METHOD_VERSION
    baseline = initial_baseline or methodology_changed
    state_rows = {}
    for offer in listings:
        old = prior.get(offer["id"], {}) if not initial_baseline else {}
        first_seen = old.get("first_seen_at") or offer["observed_at"]
        offer["first_seen_at"] = first_seen
        previous_price = old.get("price_eur")
        same_basis = old.get("vat_basis") == offer["vat_basis"] and old.get("price_currency") == offer["price_currency"] and bool(old.get("price_total_derived", False)) == offer["price_total_derived"]
        area_changed = bool(old) and "area_sqm" in old and old.get("area_sqm") != offer["area_sqm"]
        basis_changed = bool(old) and not same_basis
        comparison_basis = digest(json_bytes({"cohort": cohort_key(offer), "condition": offer["analysis"]["condition"], "building_year": actual_building_year(offer), "location_key": (offer.get("location") or {}).get("key")}))
        comparison_basis_changed = bool(old) and old.get("comparison_basis") != comparison_basis
        if offer["price_total_derived"]:
            # An edited area changes the calculated total without a price cut.
            same_basis = same_basis and old.get("area_sqm") == offer["area_sqm"]
            previous_quote = number(old.get("quoted_eur_per_sqm"))
            current_quote = offer["quoted_eur_per_sqm"]
        else:
            previous_quote, current_quote = previous_price, offer["price_eur"]
        drop = round((1 - current_quote / previous_quote) * 100, 2) if same_basis and previous_quote and current_quote else None
        change = "historical" if historical else "baseline" if baseline else "new" if not old else "attributes_changed" if comparison_basis_changed else "price_drop" if drop and drop > 0 else "became_eligible" if offer["comparison"]["eligible"] and not old.get("eligible") else "unchanged"
        offer["change"] = {"kind": change, "previous_price_eur": previous_price if same_basis else None, "price_drop_pct": drop if drop and drop > 0 else None, "price_basis_changed": basis_changed, "area_changed": area_changed, "comparison_basis_changed": comparison_basis_changed}
        if basis_changed or area_changed or comparison_basis_changed:
            offer["comparison"]["alert_eligible"] = False
        history = list(old.get("price_history", []))
        observation = {"observed_at": offer["observed_at"], "price_eur": offer["price_eur"], "price_currency": offer["price_currency"], "vat_basis": offer["vat_basis"], "price_total_derived": offer["price_total_derived"], "quoted_eur_per_sqm": offer["quoted_eur_per_sqm"], "area_sqm": offer["area_sqm"], "area_basis": offer["analysis"]["area_basis"]}
        if not history or any(history[-1].get(field) != observation[field] for field in ("price_eur", "price_currency", "vat_basis", "price_total_derived", "quoted_eur_per_sqm", "area_sqm", "area_basis")):
            history.append(observation)
        state_rows[offer["id"]] = {"first_seen_at": first_seen, "last_seen_at": offer["observed_at"], "price_eur": offer["price_eur"], "price_currency": offer["price_currency"], "vat_basis": offer["vat_basis"], "price_total_derived": offer["price_total_derived"], "quoted_eur_per_sqm": offer["quoted_eur_per_sqm"], "area_sqm": offer["area_sqm"], "eligible": offer["comparison"]["eligible"], "observation_count": old.get("observation_count", 0) + 1, "price_history": history[-64:]}
        state_rows[offer["id"]].update({"comparison_basis": comparison_basis, "area_basis": offer["analysis"]["area_basis"]})
    # Inactive IDs and price changes survive for 90 days. Absence is not a sale.
    if not initial_baseline:
        current_date = datetime.fromisoformat(captured_at.replace("Z", "+00:00")).date()
        for identifier, row in prior.items():
            if identifier in state_rows:
                continue
            try:
                seen = datetime.fromisoformat(row["last_seen_at"].replace("Z", "+00:00")).date()
            except (ValueError, KeyError):
                continue
            if 0 <= (current_date - seen).days <= 90:
                state_rows[identifier] = row
        for row in state_rows.values():
            bounded_history = []
            for item in row.get("price_history", []):
                try:
                    seen = datetime.fromisoformat(item["observed_at"].replace("Z", "+00:00")).date()
                except (ValueError, KeyError):
                    continue
                if 0 <= (current_date - seen).days <= 90:
                    bounded_history.append(item)
            row["price_history"] = bounded_history[-64:]
    neighborhoods = []
    by_neighborhood = defaultdict(list)
    for offer in listings:
        if offer["neighborhood_key"]:
            by_neighborhood[offer["neighborhood_key"]].append(offer)
    for key, offers in by_neighborhood.items():
        valid = [x for x in offers if comparable(x)]
        inclusive = [x for x in valid if x["vat_basis"] == "included"]
        cohort_groups = defaultdict(list)
        for offer in valid:
            cohort_groups[cohort_key(offer)].append(offer)
        cohorts = [{"key": cohort, "type_key": rows[0]["type_key"], "type_bg": rows[0]["type_bg"], "vat_basis": rows[0]["vat_basis"], "construction_cohort": rows[0]["construction_cohort"], **stats([x["eur_per_sqm"] for x in rows])} for cohort, rows in sorted(cohort_groups.items())]
        neighborhoods.append({"key": key, "label_bg": offers[0]["neighborhood_bg"], "count": len(offers), "priced_count": len([x for x in offers if x["price_eur"] is not None]), "comparable_count": len(valid), "headline_vat_basis": "included", **stats([x["eur_per_sqm"] for x in inclusive]), "median_price_eur": round(statistics.median([x["price_eur"] for x in inclusive]), 2) if inclusive else None, "cohorts": cohorts})
        neighborhoods[-1]["count"] = len(offers)
        neighborhoods[-1]["headline_priced_count"] = len(inclusive)
        neighborhoods[-1]["by_vat_basis"] = vat_statistics(valid)
    neighborhoods.sort(key=lambda x: x["label_bg"].casefold())
    candidate_sort = lambda x: (-x["comparison"].get("below_median_pct", 0), -x["comparison"]["peer_count"], x["id"])
    eligible = sorted([x for x in listings if x["comparison"]["eligible"]], key=candidate_sort)
    # No bootstrap or initial-live alert flood. Subsequent same-day captures
    # may report genuinely new IDs or changed prices; unchanged IDs do not.
    same_capture_date = timestamp(previous.get("captured_at"), date_only=True) == timestamp(captured_at, date_only=True)
    daily_candidates = [] if baseline else [x for x in eligible if x["change"]["kind"] in {"new", "price_drop", "became_eligible"} and not x["change"]["price_basis_changed"] and not x["change"]["area_changed"]]
    flags = Counter(flag for x in listings for flag in x["quality_flags"])
    counts = {"source_row_count": len(capture["offers"]), "valid_id_row_count": len(normalized), "listing_count": len(listings), "duplicate_row_count": len(normalized) - len(listings), "rejected_id_row_count": rejected, "priced_count": sum(x["price_eur"] is not None for x in listings), "comparable_count": sum(comparable(x) for x in listings), "eligible_count": len(eligible), "daily_candidate_count": len(daily_candidates), "quality_flags": dict(flags)}
    counts["derived_total_count"] = sum(x["price_total_derived"] for x in listings)
    counts["seller_types"] = dict(Counter(x["seller_type"] for x in listings))
    counts["location_precisions"] = dict(Counter(x["location"]["precision"] if x["location"] else "unknown" for x in listings))
    counts["located_count"] = sum(x["location"] is not None for x in listings)
    counts["description_analyzed_count"] = sum(x["analysis"]["description_analyzed"] for x in listings)
    counts["asking_estimate_count"] = sum(x["asking_estimate"]["status"] == "available" for x in listings)
    counts["detailed_estimate_count"] = sum(x["asking_estimate"]["status"] == "available" and x["asking_estimate"]["evidence"]["quality"] == "detailed" for x in listings)
    counts["construction_types"] = dict(Counter(x["property_features"]["construction_type"] for x in listings))
    counts["furnishing_types"] = dict(Counter(x["property_features"]["furnishing"] for x in listings))
    counts["agency_identity_count"] = sum(x["agency"] is not None for x in listings)
    counts["agency_phone_count"] = sum(bool(x["agency"] and x["agency"]["phone"]) for x in listings)
    counts["terms_review_count"] = sum(bool(x["analysis"]["risk_flags"]) for x in listings)
    source_dates = [x["observed_at"] for x in listings if x["observed_at"]]
    sources = capture.get("completion", {}).get("sources", [])
    coverage = {"city_bg": "София", "scope_bg": capture.get("scope_bg") or "Продажби на 2-, 3- и 4-стайни апартаменти; обяви, а не сключени сделки", "construction_year_min": capture.get("scope", {}).get("construction_year_min"), "price_min_eur": capture.get("scope", {}).get("price_min_eur"), "price_scope": capture.get("scope", {}).get("price_scope"), "source_date_min": min(source_dates), "source_date_max": max(source_dates), "captured_source_count": len(sources), "requested_neighborhood_count": capture.get("requested_neighborhood_count"), "captured_neighborhood_count": capture.get("captured_neighborhood_count"), "neighborhood_count": len(neighborhoods), "is_historical": historical, "availability_verified": False, "capture_complete": not historical, "notes_bg": "Исторически импорт от септември 2024 г. Наличността на обявите не е проверена." if historical else "Завършен обход на избраните категории и ценови диапазони. Обява, видяна при обхода, не гарантира наличност."}
    included = [x for x in listings if comparable(x) and x["vat_basis"] == "included"]
    totals = {"headline_vat_basis": "included", "headline_priced_count": len(included), **stats([x["eur_per_sqm"] for x in included]), "median_price_eur": round(statistics.median([x["price_eur"] for x in included]), 2) if included else None}
    totals.pop("count", None)
    totals["by_vat_basis"] = vat_statistics([x for x in listings if comparable(x)])
    history = list(previous.get("summary_history", [])) if not initial_baseline else []
    history_point = {"captured_at": captured_at, "listing_count": len(listings), "headline_priced_count": len(included), "median_eur_per_sqm": totals["median_eur_per_sqm"]}
    history_point["vat_medians"] = {x["vat_basis"]: x["median_eur_per_sqm"] for x in totals["by_vat_basis"]}
    if history and timestamp(history[-1]["captured_at"], date_only=True) == timestamp(captured_at, date_only=True):
        history[-1] = history_point
    else:
        history.append(history_point)
    identifier, generated = uuid4().hex, utc_now()
    type_count = Counter(x["type_key"] for x in listings if x["type_key"])
    type_labels = {x["type_key"]: x["type_bg"] for x in listings if x["type_key"]}
    vat_count = Counter(x["vat_basis"] for x in listings)
    summary = {"version": VERSION, "snapshot_id": identifier, "generated_at": generated, "captured_at": captured_at, "mode": "historical_bootstrap" if historical else "daily", "coverage": coverage, "counts": counts, "totals": totals, "neighborhoods": neighborhoods,
        "filter_options": {"neighborhoods": [{"key": x["key"], "label_bg": x["label_bg"], "count": x["count"]} for x in neighborhoods], "types": [{"key": key, "label_bg": type_labels[key], "count": count} for key, count in sorted(type_count.items())], "vat_bases": [{"key": key, "label_bg": VAT_LABELS[key], "count": count} for key, count in sorted(vat_count.items())]},
        "methodology": {"method_version": COMPARISON_METHOD_VERSION, "price_basis": "advertised_asking_price", "currency": "EUR", "bgn_per_eur": BGN_PER_EUR, "minimum_peer_count": MIN_PEERS, "deal_discount_pct": DEAL_DISCOUNT_PCT, "extreme_discount_pct": EXTREME_DISCOUNT_PCT, "peer_fields": ["neighborhood_key", "type_key", "vat_basis", "construction_cohort", "completion_stage", "area_basis", "floor_band", "construction_type", "furnishing", "elevator", "actual_year_evidence"], "area_tolerance_pct": AREA_TOLERANCE * 100, "building_year_tolerance": YEAR_TOLERANCE, "floor_tolerance": FLOOR_TOLERANCE, "condition_when_known": True, "street_sample_minimum": MIN_PEERS, "comparison_excludes_self": True, "trimmed_mean_each_tail_pct": 10, "trimmed_mean_min_count": 20,
            "notes_bg": ["Сравняваме офертни цени, а не продажни сделки или оценка на имота.", "Референтната група е същият квартал, тип апартамент, ДДС и година на строеж преди/от 2021 г. Неизвестните години и неизвестният ДДС образуват отделни групи.", "Самата обява е изключена от референтната медиана. Нужни са поне 8 други обяви.", "Годината на строеж и площта са заявени от рекламодателя. Групата с неизвестна година може да съдържа различни възрасти на сградите. Няма автоматично свързване с кадастър или официален Акт 16.", "Една обява се брои веднъж по ID. Различни обяви за един физически имот може да останат отделни.", "Средните стойности в заглавните показатели включват само цените с включен/неначисляван ДДС. Цени без или с неизвестен ДДС не се смесват с тях.", "Неизвестният ДДС позволява само сравнение с други цени с неизвестен ДДС. Тези сравнения не предполагат крайна цена и по подразбиране не изпращат известия.", "При обявена цена за кв.м общата сума се изчислява като тази цена × обявената площ и е отбелязана като изчислена. Промяна на площта или ценовата база не създава известие за намалена цена.", "Липсваща валута не се предполага. Не начисляваме ДДС автоматично."]},
        "top_offers": eligible[:30], "daily_candidates": daily_candidates[:50], "daily_candidates_total": len(daily_candidates), "history": {"baseline": baseline, "alert_baseline_at": captured_at if baseline else previous.get("alert_baseline_at"), "same_day_refresh": same_capture_date, "series": history[-180:], "notes_bg": "До 180 завършени дневни снимки. Историята на отделна обява пази до 64 промени в цената за 90 дни; отсъствието не означава продажба."}}
    summary["history"]["methodology_changed"] = methodology_changed
    summary["coverage"]["enrichment"] = capture.get("enrichment", {})
    summary["coverage"]["location_reference_date"] = resolver.reference_date
    summary["coverage"]["location_source_url"] = "https://kais.cadastre.bg/bg/OpenData"
    summary["filter_options"]["sellers"] = [{"key": key, "count": count} for key, count in sorted(counts["seller_types"].items())]
    summary["filter_options"]["precisions"] = [{"key": key, "count": count} for key, count in sorted(counts["location_precisions"].items())]
    summary["methodology"]["notes_bg"][1] = "Референтната група е същият квартал, тип, ДДС, възрастова група, заявен етап, база на площта, етажна група, конструкция, обзавеждане и асансьор. Неизвестните конструкция/обзавеждане/асансьор са отделни групи. Площите са в рамките на ±25%; при известно състояние сравняваме същото състояние. Посочените числови етажи са в рамките на ±2 етажа. Изрично посочените реални години на построяване/въвеждане са в рамките на ±10 години; планирана година не е възраст. При поне 8 подходящи обяви на същата референтна улица използваме тях, иначе квартала."
    summary["methodology"]["notes_bg"].extend([
        "Ориентирът е медиана на сравними офертни цени × заявената площ. Диапазонът е между 25-ия и 75-ия персентил на тези обяви; не е доверителен интервал или експертна оценка на пазарна стойност.",
        "Текстовите признаци са заявления на рекламодателя. Дялове, търгове, наематели, нежилищен статут и задължителни доплащания изискват проверка и изключват автоматичното определяне като възможност.",
        "Показаното местоположение е обобщена референтна зона на улица или квартал от кадастралните отворени данни. Не е местоположение на конкретен апартамент. Не използваме адреса на офиса на агенцията.",
        "Профил на агенция означава брокер/агенция; частен продавач се отбелязва само при изрично заявление. Липсата на профил не доказва частен продавач."
        ,"Позицията под/в/над диапазона се определя спрямо 25-ия и 75-ия персентил на сравнимите оферти. Липсващите характеристики и ширината на диапазона остават видими; няма произволни процентни корекции или обещание за продажна цена. Различни обяви за един имот може да останат в извадката."
    ])
    audit = dict(capture.get("audit", {}))
    full_completion = capture.get("completion", {})
    if not historical:
        audit["completion"] = full_completion
    state = {"version": VERSION, "snapshot_id": identifier, "mode": summary["mode"], "captured_at": captured_at, "comparison_method_version": COMPARISON_METHOD_VERSION, "alert_baseline_at": summary["history"]["alert_baseline_at"], "listings": state_rows, "summary_history": history[-180:], "audit": audit, "enrichment_cache": capture.get("enrichment_cache", {})}
    root = Path(output)
    folder = root / "market-snapshots" / identifier
    folder.mkdir(parents=True, exist_ok=False)
    def write_file(name, value, limit):
        payload = json_bytes(value)
        if len(payload) > limit:
            raise ValueError(f"Prepared {name} exceeds its publication limit")
        (folder / name).write_bytes(payload)
        return {"path": name, "bytes": len(payload), "sha256": digest(payload)}
    summary_desc = write_file("summary.json", summary, MAX_SUMMARY_BYTES)
    state_payload = json_bytes(state)
    if len(state_payload) > MAX_STATE_BYTES:
        raise ValueError("Expanded private state exceeds its byte limit")
    compressed_buffer = io.BytesIO()
    with gzip.GzipFile(filename="", mode="wb", fileobj=compressed_buffer, compresslevel=6, mtime=0) as stream:
        stream.write(state_payload)
    compressed_state = compressed_buffer.getvalue()
    if len(compressed_state) > MAX_STATE_COMPRESSED_BYTES:
        raise ValueError("Compressed private state exceeds the hosting-safe byte limit")
    (folder / "state.json.gz").write_bytes(compressed_state)
    state_desc = {"path": "state.json.gz", "bytes": len(compressed_state), "sha256": digest(compressed_state), "encoding": "gzip", "expanded_bytes": len(state_payload), "expanded_sha256": digest(state_payload)}
    shards, rows = [], []
    empty_size = len(json_bytes({"snapshot_id": identifier, "listings": []}))
    shard_size = empty_size
    def close_shard(rows):
        if len(shards) >= MAX_SHARDS:
            raise ValueError("Market dataset exceeds its shard count limit")
        desc = write_file(f"listings-{len(shards):03}.json", {"snapshot_id": identifier, "listings": rows}, MAX_SHARD_BYTES)
        desc["count"] = len(rows)
        shards.append(desc)
    for offer in listings:
        offer_size = len(json_bytes(offer)) - 1
        next_size = shard_size + offer_size + (1 if rows else 0)
        if rows and (len(rows) >= SHARD_ROWS or next_size > MAX_SHARD_BYTES):
            close_shard(rows)
            rows, shard_size = [], empty_size
        rows.append(offer)
        shard_size += offer_size + (1 if len(rows) > 1 else 0)
    if rows:
        close_shard(rows)
    capture_meta = {key: capture[key] for key in ("source", "started_at", "completed_at", "complete", "scope", "scope_bg", "completion", "requested_neighborhood_count", "captured_neighborhood_count") if key in capture}
    if not historical:
        # Leaf URL/page-ID evidence can exceed 256 KiB for all catalogue areas.
        # Keep it privately in state; the manifest commits to its exact bytes.
        expected = sum(x["expected_count"] for x in sources)
        captured = sum(x["captured_count"] for x in sources)
        capture_meta["completion"] = {
            "source_count": len(sources), "completed_source_count": len(sources),
            "expected_row_count": expected, "captured_row_count": captured,
            "total_unique": len(listings), "evidence_sha256": digest(json_bytes(full_completion)),
            "evidence_location": "state.audit.completion",
        }
    capture_meta["source_url"] = capture.get("source_url") or "https://www.imot.bg/"
    manifest = {"version": VERSION, "snapshot_id": identifier, "generated_at": generated, "captured_at": captured_at, "mode": summary["mode"], "capture": capture_meta, "coverage": coverage, "counts": counts, "summary_file": summary_desc, "state_file": state_desc, "listing_shards": shards, "limits": {"max_listings": MAX_LISTINGS, "max_shard_bytes": MAX_SHARD_BYTES, "max_summary_bytes": MAX_SUMMARY_BYTES, "max_state_bytes": MAX_STATE_BYTES, "max_state_compressed_bytes": MAX_STATE_COMPRESSED_BYTES}}
    manifest_desc = write_file("manifest.json", manifest, MAX_MANIFEST_BYTES)
    pointer = {"version": VERSION, "snapshot_id": identifier, "manifest_path": f"market-snapshots/{identifier}/manifest.json", "manifest_sha256": manifest_desc["sha256"]}
    root.mkdir(parents=True, exist_ok=True)
    (root / "market-active.json").write_bytes(json_bytes(pointer))
    validate_snapshot(root)
    return manifest, summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--input", type=Path)
    source.add_argument("--bootstrap-workbook", type=Path)
    parser.add_argument("--previous-state", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.bootstrap_workbook:
        capture = workbook_capture(args.bootstrap_workbook)
    else:
        if args.input.stat().st_size > MAX_CAPTURE_BYTES:
            raise ValueError("Capture exceeds input bound")
        capture = json.loads(args.input.read_text(encoding="utf-8"))
    previous = None
    if args.previous_state and args.previous_state.exists():
        if args.previous_state.stat().st_size > MAX_STATE_BYTES:
            raise ValueError("Previous state exceeds input bound")
        previous = json.loads(args.previous_state.read_text(encoding="utf-8"))
    manifest, summary = build_snapshot(capture, args.output, previous)
    print(json.dumps({"snapshot_id": manifest["snapshot_id"], "mode": manifest["mode"], "counts": manifest["counts"], "source_date_min": summary["coverage"]["source_date_min"], "source_date_max": summary["coverage"]["source_date_max"], "summary_bytes": manifest["summary_file"]["bytes"], "listing_shards": len(manifest["listing_shards"])}))


if __name__ == "__main__":
    main()
