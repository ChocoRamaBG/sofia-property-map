#!/usr/bin/env python3
"""Build small gzip SQL imports for the Sofia map's MySQL database."""

from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
COORDINATES_DEFAULT = ROOT / "Coordinates" / "Translated_Coordinates" / "Translated_Sofia_Coords_Excel.xlsx"
PERMITS_DEFAULT = ROOT / "Razresheniya za Stroezh" / "Building_Permits_ListAllDataFinal20.05.2025.xlsx"
CERTIFICATES_DEFAULT = ROOT / "Udostoverenie za Vuvezhdane v Eksploataciya" / "Exploatation_ApprovalsAllDataFinal21.05.2025.xlsx"
VISAS_DEFAULT = ROOT / "Vizi za Proektirane" / "RegisterVisasOfProection Updated 12.05.xlsx"
OUTPUT_DEFAULT = Path(__file__).resolve().parents[1] / "seed" / "generated"

CADASTRAL_ID = re.compile(r"(?<![0-9])([0-9]{5}\.[0-9]{1,6}\.[0-9]{1,6})(?![0-9.])")
VALID_SOFIA = (42.0, 43.0, 22.0, 24.5)
VALUE_BATCH = 300
FILE_ROWS = 5000
IMPORT_TIME = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

def latest_export(directory: Path, pattern: str, fallback: Path) -> Path:
    exports = directory / "exports"
    candidates = list(exports.glob(pattern)) if exports.is_dir() else []
    return max(candidates, key=lambda path: path.stat().st_mtime) if candidates else fallback


SOURCES = {
    "coordinates": {
        "label_bg": "Преведени координати на имоти",
        "label_en": "Translated plot coordinates",
        "url": None,
        "snapshot_date": None,
    },
    "building_permits": {
        "label_bg": "Регистър на разрешенията за строеж",
        "label_en": "Building permits register",
        "url": "https://nag.sofia.bg/RegisterBuildingPermitsPortal/Index",
        "snapshot_date": "2025-05-20",
    },
    "exploitation_certificates": {
        "label_bg": "Регистър на удостоверенията за въвеждане в експлоатация",
        "label_en": "Certificates for exploitation register",
        "url": "https://nag.sofia.bg/RegisterCertificateForExploitationBuildings",
        "snapshot_date": "2025-05-21",
    },
    "projecting_visas": {
        "label_bg": "Регистър на визите за проектиране",
        "label_en": "Projecting visas register",
        "url": "https://nag.sofia.bg/RegisterVisasOfProection",
        "snapshot_date": "2025-05-12",
    },
}


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def sql(value: Any) -> str:
    value = clean(value)
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, float)):
        return str(value)
    text = str(value).replace("\\", "\\\\").replace("\0", "\\0").replace("\n", "\\n").replace("\r", "\\r").replace("\x1a", "\\Z").replace("'", "''")
    return "'" + text + "'"


def batched(values: list[tuple], size: int = VALUE_BATCH) -> Iterable[list[tuple]]:
    for start in range(0, len(values), size):
        yield values[start : start + size]


def write_sql_gzip(path: Path, statements: Iterable[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("wb") as raw:
        with gzip.GzipFile(filename="", mode="wb", fileobj=raw, mtime=0) as zipped:
            for statement in statements:
                zipped.write((statement.rstrip() + "\n").encode("utf-8"))


def workbook_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        workbook.close()
        return [], []
    header_values = [clean(value) for value in header]
    looks_like_indexes = all(str(value) == str(index) for index, value in enumerate(header_values) if value is not None)
    usable_header = (
        not looks_like_indexes
        and all(isinstance(value, str) and "�" not in value for value in header_values if value is not None)
    )
    if usable_header:
        labels = [str(value) if value is not None else f"Column {index + 1}" for index, value in enumerate(header_values)]
    else:
        labels = [f"Column {index + 1}" for index in range(len(header))]
    rows = iterator
    return workbook, (labels, rows)


def source_documents(source_code: str, path: Path) -> tuple[list[tuple], list[tuple]]:
    workbook, parsed = workbook_rows(path)
    labels, rows = parsed
    documents: list[tuple] = []
    links: list[tuple] = []
    filename = path.name
    for row in rows:
        cells = [clean(value) for value in row]
        if not any(value is not None for value in cells):
            continue
        fields = {labels[index]: value for index, value in enumerate(cells) if index < len(labels) and value is not None}
        serialized = json.dumps(fields, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        digest = hashlib.sha256((source_code + "\0" + serialized).encode("utf-8")).hexdigest()
        record_key = str(cells[0])[:190] if cells and cells[0] is not None else digest[:16]
        documents.append((source_code, digest, record_key, serialized, filename, IMPORT_TIME))
        for cadastral_id in sorted(set(CADASTRAL_ID.findall(serialized))):
            links.append((cadastral_id, source_code, digest))
    workbook.close()
    return documents, links


def make_metadata(source_code: str, file_path: Path, count: int, note: str) -> list[str]:
    info = SOURCES[source_code]
    matched_date = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", file_path.name)
    date_value = (
        f"{matched_date.group(3)}-{matched_date.group(2)}-{matched_date.group(1)}"
        if matched_date
        else info["snapshot_date"]
    )
    statements = [
        "INSERT INTO data_sources (code, label_bg, label_en, source_url, snapshot_file, snapshot_date) VALUES "
        f"({sql(source_code)}, {sql(info['label_bg'])}, {sql(info['label_en'])}, {sql(info['url'])}, {sql(file_path.name)}, {sql(date_value)}) "
        "ON DUPLICATE KEY UPDATE label_bg=VALUES(label_bg), label_en=VALUES(label_en), source_url=VALUES(source_url), snapshot_file=VALUES(snapshot_file), snapshot_date=VALUES(snapshot_date);"
    ]
    import_id = f"{source_code}-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    statements.append(
        "INSERT INTO data_imports (id, source_code, source_file, imported_at, records_loaded, notes) VALUES "
        f"({sql(import_id)}, {sql(source_code)}, {sql(file_path.name)}, {sql(IMPORT_TIME)}, {int(count)}, {sql(note)}) "
        "ON DUPLICATE KEY UPDATE source_file=VALUES(source_file), imported_at=VALUES(imported_at), records_loaded=VALUES(records_loaded), notes=VALUES(notes);"
    )
    return statements


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--coordinates", type=Path, default=COORDINATES_DEFAULT)
    parser.add_argument("--permits", type=Path, default=latest_export(ROOT / "Razresheniya za Stroezh", "Building_Permits_ListAllDataFinal*.xlsx", PERMITS_DEFAULT))
    parser.add_argument("--certificates", type=Path, default=latest_export(ROOT / "Udostoverenie za Vuvezhdane v Eksploataciya", "Exploatation_ApprovalsAllDataFinal*.xlsx", CERTIFICATES_DEFAULT))
    parser.add_argument("--visas", type=Path, default=latest_export(ROOT / "Vizi za Proektirane", "Projecting_Visas_*.xlsx", VISAS_DEFAULT))
    parser.add_argument("--output", type=Path, default=OUTPUT_DEFAULT)
    parser.add_argument("--documents-only", action="store_true", help="Skip the large coordinates workbook for a monthly register refresh.")
    args = parser.parse_args()

    files = {
        "coordinates": args.coordinates,
        "building_permits": args.permits,
        "exploitation_certificates": args.certificates,
        "projecting_visas": args.visas,
    }
    for code, path in files.items():
        if args.documents_only and code == "coordinates":
            continue
        if not path.is_file():
            parser.error(f"Missing {code} workbook: {path}")

    args.output.mkdir(parents=True, exist_ok=True)
    for old_file in args.output.glob("*.sql.gz"):
        old_file.unlink()

    plot_rows: dict[str, tuple[float | None, float | None, str | None]] = {}
    coordinate_count = 0
    if not args.documents_only:
        wb = load_workbook(args.coordinates, read_only=True, data_only=True)
        sheet = wb.active
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        for row in rows:
            if len(row) < 3:
                continue
            cadastral_id = clean(row[0])
            try:
                lat, lng = float(row[1]), float(row[2])
            except (TypeError, ValueError):
                continue
            if not cadastral_id or not (VALID_SOFIA[0] <= lat <= VALID_SOFIA[1] and VALID_SOFIA[2] <= lng <= VALID_SOFIA[3]):
                continue
            plot_rows[str(cadastral_id)] = (lat, lng, args.coordinates.name)
        wb.close()
        coordinate_count = len(plot_rows)

    documents_by_source: dict[str, list[tuple]] = {}
    links: list[tuple] = []
    for source_code, path in (
        ("building_permits", args.permits),
        ("exploitation_certificates", args.certificates),
        ("projecting_visas", args.visas),
    ):
        documents, source_links = source_documents(source_code, path)
        documents_by_source[source_code] = documents
        links.extend(source_links)
        for cadastral_id, _, _ in source_links:
            plot_rows.setdefault(cadastral_id, (None, None, None))

    metadata: list[str] = []
    if not args.documents_only:
        metadata.extend(make_metadata(
            "coordinates", args.coordinates, coordinate_count,
            "Original coordinate source URL/date not recorded. Imported WGS84 points only; out-of-area and sentinel coordinates were skipped. Points are not parcel boundaries."
        ))
    for source_code, path in (
        ("building_permits", args.permits),
        ("exploitation_certificates", args.certificates),
        ("projecting_visas", args.visas),
    ):
        metadata.extend(make_metadata(source_code, path, len(documents_by_source[source_code]), "Workbook snapshot only; verify records in the original municipal register."))
    write_sql_gzip(args.output / "00_sources_and_import_log.sql.gz", metadata)

    all_plots = sorted(
        (cadastral_id, lat, lng, file_name, IMPORT_TIME)
        for cadastral_id, (lat, lng, file_name) in plot_rows.items()
    )
    for part, start in enumerate(range(0, len(all_plots), FILE_ROWS), start=1):
        chunk = all_plots[start : start + FILE_ROWS]
        statements = []
        for group in batched(chunk):
            values = ",\n".join("(" + ", ".join(sql(value) for value in row) + ")" for row in group)
            statements.append(
                "INSERT INTO parcels (cadastral_id, latitude, longitude, coordinate_source_file, updated_at) VALUES\n"
                + values
                + "\nON DUPLICATE KEY UPDATE latitude=COALESCE(VALUES(latitude), latitude), longitude=COALESCE(VALUES(longitude), longitude), "
                "coordinate_source_file=COALESCE(VALUES(coordinate_source_file), coordinate_source_file), updated_at=VALUES(updated_at);"
            )
        write_sql_gzip(args.output / f"01_parcels_{part:03d}.sql.gz", statements)

    document_statements: list[str] = []
    all_documents = [record for source_docs in documents_by_source.values() for record in source_docs]
    for group in batched(all_documents):
        values = ",\n".join("(" + ", ".join(sql(value) for value in row) + ")" for row in group)
        document_statements.append(
            "INSERT INTO source_documents (source_code, document_hash, record_key, raw_fields, snapshot_file, imported_at) VALUES\n"
            + values
            + "\nON DUPLICATE KEY UPDATE record_key=VALUES(record_key), raw_fields=VALUES(raw_fields), snapshot_file=VALUES(snapshot_file), imported_at=VALUES(imported_at);"
        )
    for group in batched(sorted(set(links))):
        values = ",\n".join("(" + ", ".join(sql(value) for value in row) + ")" for row in group)
        document_statements.append(
            "INSERT IGNORE INTO parcel_documents (parcel_cadastral_id, source_code, document_hash) VALUES\n" + values + ";"
        )
    write_sql_gzip(args.output / "02_documents_and_exact_id_links.sql.gz", document_statements)

    print(f"Created SQL import files in: {args.output}")
    print(f"Coordinate points: {coordinate_count:,}" if not args.documents_only else "Coordinate workbook skipped.")
    print(f"Plot IDs including document-only IDs: {len(plot_rows):,}")
    print(f"Register documents: {sum(len(rows) for rows in documents_by_source.values()):,}")
    print(f"Exact cadastral-ID links: {len(set(links)):,}")
    print("Import files in filename order after importing public/schema.sql.")


if __name__ == "__main__":
    main()
